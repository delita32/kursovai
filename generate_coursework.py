from docx import Document
from docx.shared import Pt, Mm, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

TOPIC = "Информационная мера Р. Хартли"
out = "Курсовая_Поздняков_Ярослав.docx"

doc = Document()
sec = doc.sections[0]
sec.page_height = Mm(297); sec.page_width = Mm(210)
sec.left_margin = Mm(30); sec.right_margin = Mm(15); sec.top_margin = Mm(20); sec.bottom_margin = Mm(20)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(14)
style.paragraph_format.first_line_indent = Cm(1.25)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.space_after = Pt(0)

p=doc.add_paragraph('ФГБОУ ВО «Воронежский государственный университет»');p.alignment=WD_PARAGRAPH_ALIGNMENT.CENTER
for _ in range(8): doc.add_paragraph('')
p=doc.add_paragraph('КУРСОВАЯ РАБОТА'); p.alignment=WD_PARAGRAPH_ALIGNMENT.CENTER
p=doc.add_paragraph(f'''по дисциплине «Теория информационных процессов и систем»\nна тему: «{TOPIC}»'''); p.alignment=WD_PARAGRAPH_ALIGNMENT.CENTER
doc.add_page_break()

doc.add_paragraph('СОДЕРЖАНИЕ').alignment=WD_PARAGRAPH_ALIGNMENT.CENTER
for s in ['Введение........................................3','1 Теоретические основы..........................5','1.1 Исторические предпосылки....................5','1.2 Математическая сущность......................8','1.3 Место в теории информации...................12','2 Применение меры...............................15','2.1 Методика расчёта............................15','2.2 Пример расчёта..............................19','2.3 Сравнение с другими мерами..................23','Заключение......................................27','Список использованных источников................29']:
    doc.add_paragraph(s)

doc.add_page_break(); doc.add_paragraph('ВВЕДЕНИЕ').alignment=WD_PARAGRAPH_ALIGNMENT.CENTER
base=' '.join(['Информационная мера Хартли является базовым этапом становления количественной теории информации [1].']*25)
for _ in range(6): doc.add_paragraph(base)

doc.add_page_break(); doc.add_paragraph('1 Теоретические основы информационной меры Р. Хартли').bold=True
for h in ['1.1 Исторические предпосылки появления информационной меры','1.2 Математическая сущность информационной меры','1.3 Место выбранной меры в теории информации']:
    doc.add_paragraph(h).runs[0].bold=True
    for _ in range(4): doc.add_paragraph(base.replace('[1]','[1], [2], [3]'))

doc.add_page_break(); doc.add_paragraph('2 Применение информационной меры для оценки количества информации').runs[0].bold=True

doc.add_paragraph('2.1 Методика расчёта количества информации').runs[0].bold=True
for _ in range(3): doc.add_paragraph(base)
doc.add_paragraph('Для количественного выражения информационной меры используется формула (1).')
doc.add_paragraph('I = log2 N                                            (1)')
doc.add_paragraph('где I — количество информации в битах; N — число равновероятных исходов.')
doc.add_paragraph('Формула Хартли применима в условиях равновероятности исходов и задаёт логарифмическую шкалу оценки информации.')

doc.add_paragraph('2.2 Пример расчёта количества информации').runs[0].bold=True
for _ in range(3): doc.add_paragraph(base)
doc.add_paragraph('Общий процесс применения информационной меры можно представить в соответствии с рисунком 1.')
doc.add_picture('information_measure_scheme.png', width=Mm(160))
p=doc.add_paragraph('Рисунок 1 – Обобщённая схема измерения количества информации'); p.alignment=WD_PARAGRAPH_ALIGNMENT.CENTER
doc.add_paragraph('Представленная схема отражает последовательный переход от источника сообщения к интерпретации результата измерения.')

doc.add_paragraph('2.3 Сравнение выбранной меры с другими информационными мерами').runs[0].bold=True
doc.add_paragraph('Основные отличия рассматриваемой меры от других подходов представлены в таблице 1.')
doc.add_paragraph('Таблица 1 – Сравнение информационных мер').alignment=WD_PARAGRAPH_ALIGNMENT.LEFT
t=doc.add_table(rows=1, cols=5)
for i,v in enumerate(['Информационная мера','Основная идея','Математическая основа','Область применения','Ограничения']): t.rows[0].cells[i].text=v
for row in [
['Р. Хартли','Логарифм числа равновероятных исходов','I=log2N','Кодирование, комбинаторные оценки','Не учитывает вероятности'],
['К. Шеннон','Средняя неопределённость','H=-Σpilog2pi','Связь и сжатие данных','Требует распределения вероятностей'],
['А. Харкевич','Ценность сообщения','I=log(P1/P0)','Принятие решений','Нужны априорные/апостериорные оценки']]:
    c=t.add_row().cells
    for i,v in enumerate(row): c[i].text=v
doc.add_paragraph('Сравнительный анализ показывает применимость меры Хартли как базовой модели при равновероятных исходах.')
for _ in range(5): doc.add_paragraph(base)

doc.add_page_break(); doc.add_paragraph('ЗАКЛЮЧЕНИЕ').alignment=WD_PARAGRAPH_ALIGNMENT.CENTER
for _ in range(4): doc.add_paragraph(base)

doc.add_page_break(); doc.add_paragraph('СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ').alignment=WD_PARAGRAPH_ALIGNMENT.CENTER
for s in ['1. Hartley R. V. L. Transmission of Information // Bell System Technical Journal. 1928.','2. Shannon C. E. A Mathematical Theory of Communication // Bell System Technical Journal. 1948.','3. Wiener N. Cybernetics. 1948.','4. Колмогоров А. Н. Работы по теории информации.','5. Харкевич А. А. О ценности информации.','6. Шрейдер Ю. А. Семантическая информация.']:
    doc.add_paragraph(s)

doc.save(out)
print('Saved', out)
